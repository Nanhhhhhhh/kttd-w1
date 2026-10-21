"""
Script tạo slide PowerPoint (PPTX) chất lượng cao cho báo cáo Tuần 6 - Nhóm W1.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

DARK_BG = RGBColor(15, 23, 42)        # #0f172a
CARD_BG = RGBColor(30, 41, 59)        # #1e293b
ACCENT_BLUE = RGBColor(56, 189, 248)  # #38bdf8
TEXT_WHITE = RGBColor(248, 250, 252)
TEXT_MUTED = RGBColor(148, 163, 184)
GREEN_COLOR = RGBColor(52, 211, 153)


def add_base_slide(title_text: str):
    slide = prs.slides.add_slide(blank_layout)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK_BG
    bg.line.fill.background()

    # Header
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.9))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = "Arial"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    # Subheader / divider line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.3), Inches(11.733), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(51, 65, 85)
    line.line.fill.background()

    # Footer
    ft = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.4))
    fp = ft.text_frame.paragraphs[0]
    fp.text = "Học phần: Kiểm thử phần mềm | Nhóm W1 – Tuần 6: Nhập liệu & phím đặc biệt"
    fp.font.size = Pt(11)
    fp.font.color.rgb = TEXT_MUTED
    return slide


# Slide 1: Cover
slide1 = prs.slides.add_slide(blank_layout)
bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = DARK_BG
bg1.line.fill.background()

tb1 = slide1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(4.0))
tf1 = tb1.text_frame
tf1.word_wrap = True

p1 = tf1.paragraphs[0]
p1.text = "HỌC PHẦN KIỂM THỬ PHẦN MỀM TỰ ĐỘNG"
p1.font.size = Pt(18)
p1.font.bold = True
p1.font.color.rgb = ACCENT_BLUE

p2 = tf1.add_paragraph()
p2.text = "BÁO CÁO TUẦN 06: CHUYÊN ĐỀ TÀI LIỆU & CHỨC NĂNG 1"
p2.font.size = Pt(36)
p2.font.bold = True
p2.font.color.rgb = TEXT_WHITE
p2.space_before = Pt(14)

p3 = tf1.add_paragraph()
p3.text = (
    "Chuyên đề: Nhập liệu & phím đặc biệt (Inputs & Special Keys)\n"
    "Trang kiểm thử: the-internet.herokuapp.com (/key_presses, /inputs)\n"
    "Nhóm thực hiện: Nhóm W1 (Web Testing Team 1)\n"
    "Thời lượng báo cáo: 15 Phút"
)
p3.font.size = Pt(16)
p3.font.color.rgb = TEXT_MUTED
p3.space_before = Pt(20)

# Slide 2: Agenda
slide2 = add_base_slide("1. NỘI DUNG VÀ THỜI LƯỢNG BÁO CÁO (15 PHÚT)")
tb2 = slide2.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.0))
tf2 = tb2.text_frame
tf2.word_wrap = True

agenda_items = [
    ("1. Quy trình làm việc nhóm trên GitHub (2 phút)", "Phân nhánh Git, Quy ước commit, Quản lý tiến độ bằng GitHub Issues & Pull Requests."),
    ("2. Tóm tắt chuyên đề tài liệu Selenium-Python (6 phút)", "Mục 2.1-2.4 (unittest.TestCase, setUp, tearDown), Mục 3.1 & 7.20 (WebElement), Mục 4.1 & 7.4 (By.ID & Keys)."),
    ("3. Thiết kế Kịch bản kiểm thử & Triển khai Mã nguồn (4 phút)", "6 Test Cases trên /key_presses và /inputs, ứng dụng đủ 13 hàm bắt buộc."),
    ("4. Kết quả chạy kiểm thử & Minh chứng thực tế (2 phút)", "Báo cáo pytest 100% Passed (6/6), Báo cáo HTML trực quan và ảnh chụp màn hình."),
    ("5. Đánh giá, Bài học kinh nghiệm & Hỏi đáp Q&A (1 phút)", "Xử lý chặn reload form HTML, tối ưu hoá thời gian chạy eager.")
]

for title, desc in agenda_items:
    p_t = tf2.add_paragraph()
    p_t.text = "• " + title
    p_t.font.size = Pt(18)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_WHITE
    p_t.space_before = Pt(12)

    p_d = tf2.add_paragraph()
    p_d.text = "    " + desc
    p_d.font.size = Pt(14)
    p_d.font.color.rgb = TEXT_MUTED

# Slide 3: GitHub workflow
slide3 = add_base_slide("2. QUY TRÌNH LÀM VIỆC NHÓM TRÊN GITHUB")
tb3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.0))
tf3 = tb3.text_frame
tf3.word_wrap = True

gh_points = [
    ("• Mô hình phân nhánh chuẩn mực:", True),
    ("   - main: Bản phát hành ổn định, bảo vệ bằng branch protection rule, chỉ nhận code qua PR.", False),
    ("   - week-06: Nhánh tích hợp chung của tuần 6, tạo từ main.", False),
    ("   - t06-<chức_năng>: Các nhánh con độc lập (t06-testcases, t06-keypresses, t06-inputs, t06-slide).", False),
    ("• Quản lý tiến độ bằng GitHub Issues & Milestones:", True),
    ("   - Tạo Milestone: Tuần 6 – Nhập liệu & phím đặc biệt.", False),
    ("   - Phân chia 4 Issues: Gán nhãn week-06, code, tai-lieu, slide và gán người phụ trách.", False),
    ("• Quy trình Pull Request và Review code:", True),
    ("   - Mỗi thành viên hoàn thành mở PR vào week-06 kèm cú pháp: Closes #<id_issue>.", False),
    ("   - Reviewer kiểm tra, Approve trước khi Merge vào nhánh tuần. Cuối tuần mở PR từ week-06 vào main.", False)
]
for pt, is_bold in gh_points:
    p = tf3.add_paragraph()
    p.text = pt
    p.font.size = Pt(16)
    p.font.color.rgb = TEXT_WHITE if is_bold else TEXT_MUTED
    p.font.bold = is_bold
    p.space_before = Pt(6)

# Slide 4: Selenium unittest
slide4 = add_base_slide("3. TÀI LIỆU SELENIUM: CẤU TRÚC TEST VỚI UNITTEST (MỤC 2.1 – 2.4)")
tb4 = slide4.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.0))
tf4 = tb4.text_frame
tf4.word_wrap = True

u_points = [
    ("• unittest.TestCase (Mục 2.3):", True),
    ("   Kế thừa lớp chuẩn để tự động hóa phát hiện và chạy test cases.", False),
    ("• Phương thức setUp() (Mục 2.4):", True),
    ("   Chạy trước mỗi test case. Đảm bảo môi trường kiểm thử và WebDriver luôn sẵn sàng.", False),
    ("• Phương thức tearDown() (Mục 2.4):", True),
    ("   Chạy sau mỗi test case để dọn dẹp dữ liệu, giải phóng tài nguyên trình duyệt.", False),
    ("• Tối ưu hoá thực tế:", True),
    ("   Áp dụng setUpClass() và tearDownClass() tái sử dụng phiên Chrome, giảm thời gian chạy từ 30s xuống 5s.", False)
]
for pt, is_bold in u_points:
    p = tf4.add_paragraph()
    p.text = pt
    p.font.size = Pt(15)
    p.font.color.rgb = TEXT_WHITE if is_bold else TEXT_MUTED
    p.font.bold = is_bold
    p.space_before = Pt(8)

code_box4 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
code_box4.fill.solid()
code_box4.fill.fore_color.rgb = CARD_BG
code_box4.line.fill.background()
tf_c4 = code_box4.text_frame
tf_c4.word_wrap = True
tf_c4.margin_left = Inches(0.2)
p_code = tf_c4.paragraphs[0]
p_code.text = (
    "import unittest\n"
    "from selenium import webdriver\n\n"
    "class TestW1(unittest.TestCase):\n"
    "    @classmethod\n"
    "    def setUpClass(cls):\n"
    "        # Khởi tạo WebDriver dùng chung\n"
    "        cls.driver = webdriver.Chrome(options)\n\n"
    "    def setUp(self):\n"
    "        # Kiểm tra trước mỗi test\n"
    "        self.assertIsNotNone(self.driver)\n\n"
    "    def tearDown(self):\n"
    "        # Dọn dẹp sau mỗi test\n"
    "        pass\n\n"
    "    @classmethod\n"
    "    def tearDownClass(cls):\n"
    "        cls.driver.quit()"
)
p_code.font.name = "Courier New"
p_code.font.size = Pt(13)
p_code.font.color.rgb = GREEN_COLOR

# Slide 5: Locators and Keys
slide5 = add_base_slide("4. BỘ ĐỊNH VỊ BY ID & PHÍM ĐẶC BIỆT KEYS (MỤC 3.1, 4.1, 7.4, 7.20)")
tb5 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.0))
tf5 = tb5.text_frame
tf5.word_wrap = True

loc_items = [
    ("By.ID (Mục 4.1)", "Phương thức định vị nhanh nhất và có tính duy nhất cao nhất trên DOM. Dùng cho input id='target' và thông báo id='result'."),
    ("WebElement (Mục 3.1 & 7.20)", "Đối tượng đại diện cho thẻ HTML. Các phương thức: send_keys() để nhập liệu, clear() để làm rỗng ô nhập, get_attribute() để đọc giá trị thuộc tính HTML, và .text để lấy chữ hiển thị."),
    ("Keys.ENTER (Mục 7.4)", "Mô phỏng gõ phím Enter. Cần lưu ý xử lý chặn submit form tự nhiên để tránh reload mất DOM."),
    ("Keys.TAB (Mục 7.4)", "Mô phỏng phím Tab chuyển focus. Trang phản hồi dòng thông báo 'You entered: TAB'."),
    ("Keys.ARROW_UP (Mục 7.4)", "Mô phỏng phím mũi tên lên. Trên trang key_presses hiển thị 'You entered: UP', trên trang inputs giúp tăng giá trị số lên 1 đơn vị.")
]
for title, desc in loc_items:
    p_t = tf5.add_paragraph()
    p_t.text = "• " + title + ": " + desc
    p_t.font.size = Pt(15)
    p_t.font.color.rgb = TEXT_WHITE
    p_t.space_before = Pt(10)

# Slide 6: Test Case Table
slide6 = add_base_slide("5. BẢNG THIẾT KẾ KỊCH BẢN KIỂM THỬ (6 TEST CASES)")
table_shape = slide6.shapes.add_table(7, 6, Inches(0.8), Inches(1.7), Inches(11.733), Inches(4.8))
table = table_shape.table

headers = ["Mã TC", "Tên ca kiểm thử", "Trang", "Hàm sử dụng", "Kỳ vọng", "Kết quả"]
widths = [Inches(1.0), Inches(2.8), Inches(1.8), Inches(2.8), Inches(2.2), Inches(1.133)]
for i, w in enumerate(widths):
    table.columns[i].width = w

for i, h in enumerate(headers):
    cell = table.cell(0, i)
    cell.fill.solid()
    cell.fill.fore_color.rgb = CARD_BG
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = ACCENT_BLUE

rows_data = [
    ("TC01", "Gõ phím ENTER", "/key_presses", "get, By.ID, send_keys, Keys.ENTER, text", '"You entered: ENTER"', "PASSED"),
    ("TC02", "Gõ phím TAB", "/key_presses", "get, By.ID, send_keys, Keys.TAB, text", '"You entered: TAB"', "PASSED"),
    ("TC03", "Gõ phím ARROW_UP", "/key_presses", "get, By.ID, send_keys, Keys.ARROW_UP, text", '"You entered: UP"', "PASSED"),
    ("TC04", "Nhập số & kiểm tra", "/inputs", "get, send_keys, get_attribute", 'type="number", val="50"', "PASSED"),
    ("TC05", "Tăng số bằng phím lên", "/inputs", "send_keys, Keys.ARROW_UP, get_attribute", "val tăng từ 50 lên 51", "PASSED"),
    ("TC06", "Xoá nội dung ô nhập", "/inputs", "send_keys, clear, get_attribute", 'val trở về rỗng ""', "PASSED")
]

for row_idx, r_data in enumerate(rows_data, start=1):
    for col_idx, text in enumerate(r_data):
        cell = table.cell(row_idx, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_BG
        p = cell.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(12)
        if col_idx == 5:
            p.font.bold = True
            p.font.color.rgb = GREEN_COLOR
        else:
            p.font.color.rgb = TEXT_WHITE

# Slide 7: Execution Results and Screenshots
slide7 = add_base_slide("6. KẾT QUẢ THỰC THI & MINH CHỨNG THỰC TẾ")

tb7 = slide7.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(4.5), Inches(5.0))
tf7 = tb7.text_frame
tf7.word_wrap = True

p_res = tf7.paragraphs[0]
p_res.text = (
    "• Kết quả thực thi Pytest:\n"
    "   - 6/6 Ca kiểm thử đạt (100% Passed)\n"
    "   - Tổng thời gian chạy: 5.01 giây\n"
    "   - Xuất báo cáo: tuan-06/anh/report.html\n\n"
    "• Cơ chế đảm bảo chất lượng:\n"
    "   - So sánh chính xác text & attributes bằng assert\n"
    "   - Tự động chụp ảnh minh chứng từng bước vào thư mục tuan-06/anh/\n"
    "   - Chạy ổn định cả trên GUI và dòng lệnh CI/Headless"
)
p_res.font.size = Pt(14)
p_res.font.color.rgb = TEXT_WHITE

img_path1 = "/Users/macos/kttd-w1/tuan-06/anh/tc01_key_presses_enter.png"
img_path2 = "/Users/macos/kttd-w1/tuan-06/anh/tc04_inputs_send_keys.png"

if os.path.exists(img_path1):
    slide7.shapes.add_picture(img_path1, Inches(5.5), Inches(1.6), width=Inches(3.6))
    cap1 = slide7.shapes.add_textbox(Inches(5.5), Inches(4.0), Inches(3.6), Inches(0.4))
    cap1.text_frame.paragraphs[0].text = "Minh chứng TC01: Gõ ENTER"
    cap1.text_frame.paragraphs[0].font.size = Pt(11)
    cap1.text_frame.paragraphs[0].font.color.rgb = ACCENT_BLUE

if os.path.exists(img_path2):
    slide7.shapes.add_picture(img_path2, Inches(9.3), Inches(1.6), width=Inches(3.6))
    cap2 = slide7.shapes.add_textbox(Inches(9.3), Inches(4.0), Inches(3.6), Inches(0.4))
    cap2.text_frame.paragraphs[0].text = "Minh chứng TC04: Nhập số inputs"
    cap2.text_frame.paragraphs[0].font.size = Pt(11)
    cap2.text_frame.paragraphs[0].font.color.rgb = ACCENT_BLUE

# Slide 8: Lessons Learned and Conclusion
slide8 = add_base_slide("7. BÀI HỌC KINH NGHIỆM & KẾT LUẬN")
tb8 = slide8.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.0))
tf8 = tb8.text_frame
tf8.word_wrap = True

concl_items = [
    ("Xử lý biểu mẫu HTML tự động submit", "Khi nhấn Keys.ENTER vào ô input trong form, trình duyệt mặc định gửi request GET reload trang khiến element bị stale. Cần sử dụng execute_script chặn event.preventDefault() để bảo toàn DOM."),
    ("Tối ưu hoá hiệu năng chạy kiểm thử", "Cấu hình page_load_strategy='eager' giúp trình duyệt không bị treo do các script thống kê ngoài của trang demo. Áp dụng setUpClass tái sử dụng 1 browser duy nhất giúp tốc độ chạy tăng gấp 6 lần."),
    ("Giá trị của quy trình Git nhóm chuẩn mực", "Phân chia Issue từ đầu tuần giúp từng thành viên chủ động. Quy trình Review và Pull Request vào nhánh tuần trước khi gộp vào main loại bỏ hoàn toàn xung đột code.")
]
for title, desc in concl_items:
    p_t = tf8.add_paragraph()
    p_t.text = "• " + title
    p_t.font.size = Pt(17)
    p_t.font.bold = True
    p_t.font.color.rgb = ACCENT_BLUE
    p_t.space_before = Pt(12)

    p_d = tf8.add_paragraph()
    p_d.text = "   " + desc
    p_d.font.size = Pt(14)
    p_d.font.color.rgb = TEXT_WHITE

# Slide 9: QA
slide9 = prs.slides.add_slide(blank_layout)
bg9 = slide9.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg9.fill.solid()
bg9.fill.fore_color.rgb = DARK_BG
bg9.line.fill.background()

tb9 = slide9.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(3.5))
tf9 = tb9.text_frame
tf9.word_wrap = True

p_qa1 = tf9.paragraphs[0]
p_qa1.alignment = PP_ALIGN.CENTER
p_qa1.text = "CẢM ƠN CÔ VÀ CÁC BẠN ĐÃ LẮNG NGHE!"
p_qa1.font.size = Pt(36)
p_qa1.font.bold = True
p_qa1.font.color.rgb = ACCENT_BLUE

p_qa2 = tf9.add_paragraph()
p_qa2.alignment = PP_ALIGN.CENTER
p_qa2.text = "PHẦN HỎI ĐÁP – Q & A"
p_qa2.font.size = Pt(24)
p_qa2.font.bold = True
p_qa2.font.color.rgb = TEXT_WHITE
p_qa2.space_before = Pt(16)

p_qa3 = tf9.add_paragraph()
p_qa3.alignment = PP_ALIGN.CENTER
p_qa3.text = "Nhóm W1 – Kho mã nguồn: github.com/Nanhhhhhhh/kttd-w1"
p_qa3.font.size = Pt(16)
p_qa3.font.color.rgb = TEXT_MUTED
p_qa3.space_before = Pt(20)

output_pptx = "/Users/macos/kttd-w1/tuan-06/slide_tuan06_w1.pptx"
prs.save(output_pptx)
print("Saved PPTX successfully to:", output_pptx)
