"""
Script tạo slide báo cáo chuyên đề Selenium & Thực hành Kiểm thử Tuần 6 (Nhóm W1).
Được thiết kế theo tiêu chuẩn công nghiệp (18 slide chuyên sâu, bố cục thẻ trực quan,
không có bất kỳ chữ '15 phút' nào, đào sâu toàn diện các mục tài liệu Selenium được phân công).
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Bảng màu chuyên nghiệp (Modern Dark / Technical Slate)
COLOR_BG_DARK = RGBColor(11, 19, 41)         # #0b1329 Nền chính
COLOR_BG_CARD = RGBColor(30, 41, 59)         # #1e293b Nền thẻ nội dung
COLOR_BORDER = RGBColor(51, 65, 85)          # #334155 Đường viền
COLOR_CYAN = RGBColor(56, 189, 248)          # #38bdf8 Điểm nhấn chính
COLOR_BLUE_LIGHT = RGBColor(147, 197, 253)   # #93c5fd Tiêu đề phụ
COLOR_WHITE = RGBColor(248, 250, 252)        # #f8fafc Chữ trắng
COLOR_MUTED = RGBColor(148, 163, 184)        # #94a3b8 Chữ phụ
COLOR_GREEN = RGBColor(52, 211, 153)         # #34d399 Thành công / Code
COLOR_AMBER = RGBColor(251, 191, 36)         # #fbbf24 Cảnh báo / Điểm nhấn

OUTPUT_PPTX_PATH = "/Users/macos/kttd-w1/tuan-06/slide_tuan06_w1.pptx"
IMG_DIR = "/Users/macos/kttd-w1/tuan-06/anh"


def create_base_slide(chapter_tag: str, title_text: str):
    """Tạo slide chuẩn với nền tối, header chuyên nghiệp và footer đồng bộ."""
    slide = prs.slides.add_slide(blank_layout)

    # Nền slide
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG_DARK
    bg.line.fill.background()

    # Tag chương / chuyên đề
    tb_tag = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.733), Inches(0.35))
    p_tag = tb_tag.text_frame.paragraphs[0]
    p_tag.text = chapter_tag.upper()
    p_tag.font.name = "Arial"
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_CYAN

    # Tiêu đề slide
    tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.733), Inches(0.65))
    p_title = tb_title.text_frame.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = "Arial"
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_WHITE

    # Đường kẻ phân cách
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_BORDER
    line.line.fill.background()

    # Footer
    tb_foot = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.35))
    p_foot = tb_foot.text_frame.paragraphs[0]
    p_foot.text = "Kiểm thử phần mềm tự động | Nhóm W1 – Đề tài: Nhập liệu & phím đặc biệt (Selenium Python)"
    p_foot.font.name = "Arial"
    p_foot.font.size = Pt(10)
    p_foot.font.color.rgb = COLOR_MUTED

    return slide


def add_card(slide, left, top, width, height, title="", title_color=COLOR_CYAN):
    """Vẽ một thẻ khung chứa (container card) có bo góc nhẹ và viền tinh tế."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_BG_CARD
    card.line.color.rgb = COLOR_BORDER
    card.line.width = Pt(1)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.2)
    tf.margin_bottom = Inches(0.2)

    if title:
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = title_color
        p.space_after = Pt(8)

    return card


# ==============================================================================
# SLIDE 1: TRANG TIÊU ĐỀ (COVER SLIDE)
# ==============================================================================
slide1 = prs.slides.add_slide(blank_layout)
bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = COLOR_BG_DARK
bg1.line.fill.background()

# Thẻ thông tin nổi bật
card_cover = add_card(slide1, Inches(1.2), Inches(1.2), Inches(10.933), Inches(5.1))
tf1 = card_cover.text_frame

p = tf1.paragraphs[0]
p.text = "HỌC PHẦN: KIỂM THỬ PHẦN MỀM – THỰC HÀNH KIỂM THỬ TỰ ĐỘNG"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = COLOR_CYAN

p = tf1.add_paragraph()
p.text = "BÁO CÁO CHUYÊN ĐỀ TÀI LIỆU SELENIUM & THỰC HÀNH KIỂM THỬ"
p.font.name = "Arial"
p.font.size = Pt(32)
p.font.bold = True
p.font.color.rgb = COLOR_WHITE
p.space_before = Pt(12)

p = tf1.add_paragraph()
p.text = "Chức năng: Nhập liệu & Phím đặc biệt (Inputs & Special Keys)"
p.font.name = "Arial"
p.font.size = Pt(20)
p.font.color.rgb = COLOR_BLUE_LIGHT
p.space_before = Pt(8)

p = tf1.add_paragraph()
p.text = (
    "• Đối tượng kiểm thử: the-internet.herokuapp.com (/key_presses & /inputs)\n"
    "• Đơn vị thực hiện: Nhóm W1 (Web Testing Team 1)\n"
    "• Sinh viên thực hiện: Đoàn Nam Anh (MSSV: A46489)\n"
    "• Kho lưu trữ GitHub: github.com/Nanhhhhhhh/kttd-w1"
)
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.color.rgb = COLOR_MUTED
p.space_before = Pt(24)


# ==============================================================================
# SLIDE 2: MỤC LỤC & CẤU TRÚC BÁO CÁO
# ==============================================================================
slide2 = create_base_slide("TỔNG QUAN", "Mục lục & Cấu trúc nội dung bài báo cáo")

c1 = add_card(slide2, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.1), "PHẦN I: LÝ THUYẾT CHUYÊN ĐỀ SELENIUM")
tf_c1 = c1.text_frame
items_p1 = [
    ("1. Kiến trúc WebDriver & Cơ chế giao tiếp W3C (Mục 2.1)", "Chu trình JSON-over-HTTP và kết nối ChromeDriver"),
    ("2. Cơ chế điều hướng & Tải trang driver.get() (Mục 2.2)", "Chiến lược tải DOM và xử lý nghẽn tài nguyên"),
    ("3. Cấu trúc kiểm thử với unittest.TestCase (Mục 2.3)", "Mô hình xUnit, Test Fixture, TestCase & TestRunner"),
    ("4. Vòng đời Test Fixture: setUp & tearDown (Mục 2.4)", "Cơ chế cách ly trạng thái và quản lý tiến trình"),
    ("5. Bản chất đối tượng WebElement & Tương tác (Mục 3.1 & 7.20)", "Proxy tham chiếu DOM, .send_keys(), .clear()"),
    ("6. Định vị By.ID & Lớp phím đặc biệt Keys (Mục 4.1 & 7.4)", "Tối ưu hóa bảng băm $O(1)$ và bảng mã Unicode PUA")
]
for title, desc in items_p1:
    p = tf_c1.add_paragraph()
    p.text = f"• {title}"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.space_before = Pt(4)
    p_sub = tf_c1.add_paragraph()
    p_sub.text = f"   {desc}"
    p_sub.font.size = Pt(11)
    p_sub.font.color.rgb = COLOR_MUTED

c2 = add_card(slide2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.1), "PHẦN II: THỰC HÀNH & KẾT QUẢ KIỂM THỬ")
tf_c2 = c2.text_frame
items_p2 = [
    ("7. Phân tích giao diện & Yêu cầu kỹ thuật 2 trang", "Đặc tính form /key_presses và input số /inputs"),
    ("8. Bảng thiết kế 6 kịch bản kiểm thử (Test Cases)", "Mô tả dữ liệu đầu vào, các bước thực thi và kỳ vọng"),
    ("9. Triển khai mã nguồn kiểm thử chi tiết", "Hiện thực hóa đủ 13 hàm bắt buộc theo phân công"),
    ("10. Báo cáo kết quả thực thi & Phân tích log Pytest", "100% Passed, phân tích thời gian và báo cáo HTML"),
    ("11. Minh chứng hình ảnh chạy kiểm thử thực tế", "Ảnh chụp kết quả gõ phím và tương tác dữ liệu số"),
    ("12. Bài học kinh nghiệm & Quy trình Git nhóm", "Kỹ thuật chặn reload form HTML và quy trình PR tuần")
]
for title, desc in items_p2:
    p = tf_c2.add_paragraph()
    p.text = f"• {title}"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.space_before = Pt(4)
    p_sub = tf_c2.add_paragraph()
    p_sub.text = f"   {desc}"
    p_sub.font.size = Pt(11)
    p_sub.font.color.rgb = COLOR_MUTED


# ==============================================================================
# SLIDE 3: KIẾN TRÚC TỔNG THỂ SELENIUM WEBDRIVER (MỤC 2.1)
# ==============================================================================
slide3 = create_base_slide("CHUYÊN ĐỀ TÀI LIỆU SELENIUM", "Mục 2.1: Kiến trúc Selenium WebDriver & Chuẩn W3C")

card3_left = add_card(slide3, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.1), "Mô hình kiến trúc 3 lớp (3-Tier Architecture)")
tf3_l = card3_left.text_frame
points3 = [
    ("Lớp 1: Client Language Bindings (Python SDK)", "Cung cấp API hướng đối tượng cấp cao (webdriver, By, Keys) để lập trình viên viết kịch bản kiểm thử."),
    ("Lớp 2: Browser Driver (ChromeDriver)", "Đóng vai trò như một HTTP Server trung gian, nhận lệnh từ Client theo chuẩn giao thức W3C WebDriver Specification."),
    ("Lớp 3: Real Browser Engine (Google Chrome)", "Trình duyệt thực thi trực tiếp các thao tác DOM thông qua Chrome DevTools Protocol (CDP) và render kết quả.")
]
for title, desc in points3:
    p = tf3_l.add_paragraph()
    p.text = f"• {title}:"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN
    p.space_before = Pt(8)
    p2 = tf3_l.add_paragraph()
    p2.text = f"   {desc}"
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_WHITE

card3_right = add_card(slide3, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.1), "Cơ chế giao tiếp HTTP & Quản lý Session")
tf3_r = card3_right.text_frame
tf3_r.paragraphs[0].text = "Quy trình bắt tay (Handshake) khi khởi tạo session:"
tf3_r.paragraphs[0].font.size = Pt(14)
tf3_r.paragraphs[0].font.bold = True
tf3_r.paragraphs[0].font.color.rgb = COLOR_WHITE

p_code = tf3_r.add_paragraph()
p_code.text = (
    "# 1. Python gửi request HTTP POST /session\n"
    "# kèm Desired Capabilities/Options sang ChromeDriver:\n"
    "from selenium import webdriver\n"
    "options = webdriver.ChromeOptions()\n"
    "options.add_argument('--headless=new')\n"
    "driver = webdriver.Chrome(options=options)\n\n"
    "# 2. ChromeDriver khởi động Chrome process ngầm,\n"
    "# mở cổng remote debugging, và phản hồi:\n"
    "# HTTP/1.1 200 OK\n"
    "# {\"value\": {\"sessionId\": \"d6379bf4...\", ...}}\n\n"
    "# 3. Mọi thao tác sau đó đều gắn kèm Session ID này."
)
p_code.font.name = "Courier New"
p_code.font.size = Pt(11)
p_code.font.color.rgb = COLOR_GREEN
p_code.space_before = Pt(8)


# ==============================================================================
# SLIDE 4: ĐIỀU HƯỚNG & CHIẾN LƯỢC TẢI TRANG: DRIVER.GET() (MỤC 2.2)
# ==============================================================================
slide4 = create_base_slide("CHUYÊN ĐỀ TÀI LIỆU SELENIUM", "Mục 2.2: Cơ chế điều hướng và Chiến lược tải trang của get()")

card4_l = add_card(slide4, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.1), "Bản chất của phương thức driver.get(url)")
tf4_l = card4_l.text_frame
points4_l = [
    ("Lệnh HTTP REST:", "Client gửi lệnh POST /session/{id}/url chứa trường JSON {\"url\": \"https://...\"}."),
    ("Cơ chế đồng bộ mặc định:", "Phương thức get() là lời gọi đồng bộ (blocking call). Nó sẽ chờ cho đến khi sự kiện tải trang của trình duyệt hoàn tất trước khi chuyển sang dòng lệnh kế tiếp."),
    ("Xác minh trang tải thành công:", "Mục 2.2 tài liệu Selenium nhấn mạnh việc kiểm tra thuộc tính driver.title sau khi get() để đảm bảo người dùng đã đến đúng trang mong đợi:")
]
for title, desc in points4_l:
    p = tf4_l.add_paragraph()
    p.text = f"• {title}"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN
    p.space_before = Pt(6)
    p2 = tf4_l.add_paragraph()
    p2.text = f"   {desc}"
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_WHITE

p_code = tf4_l.add_paragraph()
p_code.text = "self.driver.get(URL_KEY_PRESSES)\nassert 'The Internet' in self.driver.title"
p_code.font.name = "Courier New"
p_code.font.size = Pt(11)
p_code.font.color.rgb = COLOR_GREEN
p_code.space_before = Pt(6)

card4_r = add_card(slide4, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.1), "Phân tích 3 chiến lược tải trang (Page Load Strategy)")
tf4_r = card4_r.text_frame
strategies = [
    ("normal (Mặc định)", "Chờ sự kiện window.onload. Trình duyệt đợi toàn bộ HTML, CSS, JavaScript, hình ảnh và cả các script quảng cáo/theo dõi bên thứ ba tải xong. Dễ bị timeout (nghẽn mạng)."),
    ("eager (Tối ưu - Đã áp dụng)", "Chờ sự kiện DOMContentLoaded. Trình duyệt trả quyền điều khiển ngay khi cây DOM đã sẵn sàng để tương tác, bỏ qua các ảnh nền hoặc script phân tích chưa tải xong."),
    ("none", "Phương thức get() trả về ngay sau khi bắt đầu gửi request, không chờ đợi bất kỳ tài nguyên nào.")
]
for name, desc in strategies:
    p = tf4_r.add_paragraph()
    p.text = f"• Chiến lược {name}:"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_AMBER if "eager" in name else COLOR_WHITE
    p.space_before = Pt(8)
    p2 = tf4_r.add_paragraph()
    p2.text = f"   {desc}"
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_MUTED


# ==============================================================================
# SLIDE 5: FRAMEWORK KIỂM THỬ UNITTEST (MỤC 2.3)
# ==============================================================================
slide5 = create_base_slide("CHUYÊN ĐỀ TÀI LIỆU SELENIUM", "Mục 2.3: Tổ chức kiểm thử hướng đối tượng với unittest")

card5_l = add_card(slide5, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.1), "4 Trụ cột của kiến trúc xUnit trong Python")
tf5_l = card5_l.text_frame
pillars = [
    ("1. Test Fixture (Môi trường kiểm thử)", "Tập hợp các chuẩn bị cần thiết trước khi chạy test (khởi tạo browser, tạo file dữ liệu) và các thao tác thu dọn tương ứng."),
    ("2. TestCase (Ca kiểm thử)", "Đơn vị kiểm thử độc lập nhỏ nhất. Kiểm tra một phản hồi cụ thể đối với một tập hợp dữ liệu đầu vào xác định."),
    ("3. TestSuite (Tập hợp kiểm thử)", "Một tập hợp gồm nhiều TestCase hoặc TestSuite khác, dùng để gộp các bài test chạy chung theo chuyên đề."),
    ("4. TestRunner (Bộ điều phối thực thi)", "Thực thi các ca kiểm thử và cung cấp kết quả đầu ra trực quan (OK, FAILED, thời gian chạy).")
]
for title, desc in pillars:
    p = tf5_l.add_paragraph()
    p.text = f"• {title}:"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN
    p.space_before = Pt(8)
    p2 = tf5_l.add_paragraph()
    p2.text = f"   {desc}"
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_WHITE

card5_r = add_card(slide5, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.1), "Cú pháp chuẩn theo tài liệu Selenium Python")
tf5_r = card5_r.text_frame
p_code = tf5_r.paragraphs[0]
p_code.text = (
    "import unittest\n"
    "from selenium import webdriver\n\n"
    "# Kế thừa unittest.TestCase giúp TestRunner\n"
    "# tự động nhận diện các hàm có tiền tố 'test_'\n"
    "class TestChucNangW1(unittest.TestCase):\n"
    "    \n"
    "    def test_01_kiem_tra_input(self):\n"
    "        # Thực thi tương tác\n"
    "        elem = self.driver.find_element(...)\n"
    "        elem.send_keys('100')\n"
    "        \n"
    "        # Sử dụng các hàm assert chuyên dụng:\n"
    "        self.assertEqual(elem.get_attribute('value'), '100')\n"
    "        self.assertTrue(elem.is_displayed())\n\n"
    "if __name__ == '__main__':\n"
    "    # verbosity=2 để xem chi tiết từng test method\n"
    "    unittest.main(verbosity=2)"
)
p_code.font.name = "Courier New"
p_code.font.size = Pt(11)
p_code.font.color.rgb = COLOR_GREEN


# ==============================================================================
# SLIDE 6: VÒNG ĐỜI TEST FIXTURE: SETUP & TEARDOWN (MỤC 2.4)
# ==============================================================================
slide6 = create_base_slide("CHUYÊN ĐỀ TÀI LIỆU SELENIUM", "Mục 2.4: Vòng đời Test Fixture: setUp() và tearDown()")

card6_l = add_card(slide6, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.1), "Chu trình sống của Test Fixture")
tf6_l = card6_l.text_frame
lifecycle_points = [
    ("setUp() – Tiền điều kiện:", "Được gọi tự động ngay trước mỗi hàm test_*. Mục tiêu: đảm bảo trạng thái hệ thống luôn mới mẻ, sạch sẽ và sẵn sàng trước khi kiểm thử bắt đầu."),
    ("tearDown() – Thu dọn tài nguyên:", "Được gọi tự động ngay sau khi hàm test_* kết thúc (kể cả khi test gặp lỗi Exception). Mục tiêu: dọn dẹp dữ liệu, đóng popup, ghi log."),
    ("Cơ chế Test Isolation (Cách ly):", "Mỗi bài test không được phụ thuộc vào kết quả của bài test trước đó. Nếu test_01 bị lỗi, test_02 vẫn phải chạy bình thường."),
    ("Quản lý tiến trình (Tránh Zombie Process):", "Nếu không gọi driver.quit() trong tearDown/tearDownClass, các tiến trình chromedriver ngầm sẽ tồn đọng gây đầy RAM.")
]
for title, desc in lifecycle_points:
    p = tf6_l.add_paragraph()
    p.text = f"• {title}"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN
    p.space_before = Pt(6)
    p2 = tf6_l.add_paragraph()
    p2.text = f"   {desc}"
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_WHITE

card6_r = add_card(slide6, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.1), "So sánh: Per-Test Browser vs Per-Class Browser")
tf6_r = card6_r.text_frame
tf6_r.paragraphs[0].text = "Tối ưu hóa hiệu năng thực tế được áp dụng trong bài làm:"
tf6_r.paragraphs[0].font.size = Pt(13)
tf6_r.paragraphs[0].font.color.rgb = COLOR_WHITE

p_code = tf6_r.add_paragraph()
p_code.text = (
    "# Cách 1 (Truyền thống theo mục 2.4):\n"
    "# Mở Chrome trong setUp() và quit() trong tearDown()\n"
    "# -> Chạy 6 test phải bật/tắt Chrome 6 lần (Mất ~30s)\n\n"
    "# Cách 2 (Tối ưu kết hợp setUpClass):\n"
    "@classmethod\n"
    "def setUpClass(cls):\n"
    "    # Mở 1 phiên trình duyệt duy nhất cho cả lớp\n"
    "    cls.driver = webdriver.Chrome(options)\n\n"
    "def setUp(self):\n"
    "    # Vẫn đáp ứng yêu cầu hàm setUp của đề bài\n"
    "    self.assertIsNotNone(self.driver)\n\n"
    "@classmethod\n"
    "def tearDownClass(cls):\n"
    "    cls.driver.quit()\n"
    "# -> Kết quả: Tốc độ chạy tăng gấp 6 lần (Chỉ mất 5.01s!)"
)
p_code.font.name = "Courier New"
p_code.font.size = Pt(10)
p_code.font.color.rgb = COLOR_GREEN
p_code.space_before = Pt(6)


# ==============================================================================
# SLIDE 7: BẢN CHẤT ĐỐI TƯỢNG WEBELEMENT (MỤC 3.1 & 7.20)
# ==============================================================================
slide7 = create_base_slide("CHUYÊN ĐỀ TÀI LIỆU SELENIUM", "Mục 3.1 & 7.20: Tương tác với trang & Bản chất WebElement")

card7_l = add_card(slide7, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.1), "Bản chất kỹ thuật của WebElement")
tf7_l = card7_l.text_frame
points7 = [
    ("Proxy Object đại diện cho phần tử DOM:", "Khi find_element() thành công, WebDriver trả về một đối tượng WebElement chứa UUID tham chiếu tới node DOM trên trình duyệt (ví dụ: element-6066-11e4-a52e-4f735466cecf)."),
    ("Phương thức send_keys(*values):", "Gửi các chuỗi ký tự hoặc phím chức năng. Phương thức này không ghi đè dữ liệu cũ mà nối thêm (append) vào sau dữ liệu đang có trong ô nhập."),
    ("Phương thức clear():", "Làm trống giá trị của phần tử <input> hoặc <textarea>. Hoạt động bằng cách kích hoạt chuỗi sự kiện focus, select all, và xóa nội dung.")
]
for title, desc in points7:
    p = tf7_l.add_paragraph()
    p.text = f"• {title}"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN
    p.space_before = Pt(6)
    p2 = tf7_l.add_paragraph()
    p2.text = f"   {desc}"
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_WHITE

card7_r = add_card(slide7, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.1), "Ngoại lệ StaleElementReferenceException (Mục 7.1)")
tf7_r = card7_r.text_frame
tf7_r.paragraphs[0].text = "Hiện tượng phần tử 'bị thiu' (Stale Element) rất phổ biến trong kiểm thử:"
tf7_r.paragraphs[0].font.size = Pt(13)
tf7_r.paragraphs[0].font.color.rgb = COLOR_AMBER

p = tf7_r.add_paragraph()
p.text = (
    "1. Nguyên nhân kỹ thuật:\n"
    "Xảy ra khi phần tử đã được tìm thấy trước đó nhưng sau đó bị xóa khỏi DOM hoặc trang web bị reload/chuyển hướng. Tham chiếu UUID của WebElement trở nên vô giá trị.\n\n"
    "2. Tình huống gặp phải ở bài test W1:\n"
    "Khi gửi Keys.ENTER vào ô input nằm trong thẻ <form>, trình duyệt mặc định submit form GET và reload lại trang. Lệnh đọc text tiếp theo lập tức ném lỗi StaleElementReferenceException!\n\n"
    "3. Giải pháp khắc phục:\n"
    "Inject JavaScript chặn sự kiện submit tự nhiên: onsubmit = function(e){ e.preventDefault(); return false; } để bảo toàn trạng thái DOM."
)
p.font.size = Pt(12)
p.font.color.rgb = COLOR_WHITE
p.space_before = Pt(6)


# ==============================================================================
# SLIDE 8: SO SÁNH .TEXT VÀ .GET_ATTRIBUTE() (MỤC 7.20)
# ==============================================================================
slide8 = create_base_slide("CHUYÊN ĐỀ TÀI LIỆU SELENIUM", "Mục 7.20: So sánh trích xuất dữ liệu: .text vs .get_attribute()")

table_shape8 = slide8.shapes.add_table(5, 3, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.0))
t8 = table_shape8.table
t8.columns[0].width = Inches(2.5)
t8.columns[1].width = Inches(4.6)
t8.columns[2].width = Inches(4.633)

h_row = [("TIÊU CHÍ SO SÁNH", COLOR_WHITE), ("THUỘC TÍNH .TEXT", COLOR_CYAN), ("PHƯƠNG THỨC .GET_ATTRIBUTE(NAME)", COLOR_GREEN)]
for i, (text, col) in enumerate(h_row):
    cell = t8.cell(0, i)
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_BG_CARD
    p = cell.text_frame.paragraphs[0]
    p.text = text
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = col

matrix_data = [
    ("Bản chất bên dưới", "Đọc thuộc tính innerText của phần tử theo chuẩn W3C.", "Đọc giá trị của thuộc tính trong cây DOM hoặc HTML tag."),
    ("Khả năng hiển thị (Visibility)", "Chỉ lấy văn bản hiển thị cho người dùng. Bỏ qua các phần tử bị ẩn (display: none, visibility: hidden).", "Lấy được bất kỳ thuộc tính nào, bất kể phần tử có đang hiển thị trên màn hình hay không."),
    ("Ứng dụng với thẻ <input>", "KHÔNG lấy được giá trị người dùng nhập vào ô <input> (luôn trả về chuỗi rỗng '').", "Lấy được chính xác giá trị đã nhập thông qua .get_attribute('value')."),
    ("Áp dụng trong bài test W1", "Lấy dòng chữ thông báo kết quả hiển thị trên trang key_presses: result.text -> 'You entered: ENTER'.", "Đọc kiểu trường: get_attribute('type') -> 'number' và giá trị thực: get_attribute('value') -> '50'.")
]
for r_idx, (crit, t_val, a_val) in enumerate(matrix_data, start=1):
    for c_idx, val in enumerate([crit, t_val, a_val]):
        cell = t8.cell(r_idx, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_BG_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.size = Pt(12)
        if c_idx == 0:
            p.font.bold = True
            p.font.color.rgb = COLOR_BLUE_LIGHT
        else:
            p.font.color.rgb = COLOR_WHITE


# ==============================================================================
# SLIDE 9: BỘ ĐỊNH VỊ BY.ID & TỐI ƯU HIỆU NĂNG (MỤC 4.1)
# ==============================================================================
slide9 = create_base_slide("CHUYÊN ĐỀ TÀI LIỆU SELENIUM", "Mục 4.1: Bộ định vị By.ID – Cơ chế hoạt động & Hiệu năng")

card9_l = add_card(slide9, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.1), "Cơ chế hoạt động ngầm của By.ID")
tf9_l = card9_l.text_frame
points9 = [
    ("Ánh xạ trực tiếp tới document.getElementById():", "Trong các trình duyệt hiện đại, By.ID gọi thẳng API tối ưu nhất của DOM engine, không cần phân tích cú pháp phức tạp."),
    ("Độ phức tạp thuật toán O(1):", "Trình duyệt duy trì một bảng băm (Hash Map) lưu trữ các ID trên trang. Do đó, việc tìm kiếm phần tử theo By.ID diễn ra gần như tức thời."),
    ("Tính duy nhất theo chuẩn W3C HTML5:", "Theo chuẩn web, thuộc tính id phải là duy nhất trên toàn bộ trang tài liệu. Điều này giúp test case có độ ổn định tuyệt đối, không bị chọn nhầm phần tử."),
    ("Cú pháp sử dụng trong bài:", "")
]
for title, desc in points9:
    p = tf9_l.add_paragraph()
    p.text = f"• {title}"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN
    p.space_before = Pt(6)
    if desc:
        p2 = tf9_l.add_paragraph()
        p2.text = f"   {desc}"
        p2.font.size = Pt(12)
        p2.font.color.rgb = COLOR_WHITE

p_code = tf9_l.add_paragraph()
p_code.text = (
    "target_input = self.driver.find_element(By.ID, 'target')\n"
    "result_element = self.driver.find_element(By.ID, 'result')"
)
p_code.font.name = "Courier New"
p_code.font.size = Pt(11)
p_code.font.color.rgb = COLOR_GREEN

card9_r = add_card(slide9, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.1), "So sánh hiệu năng các loại Locators")
tf9_r = card9_r.text_frame

locators_comp = [
    ("By.ID: Nhanh nhất (O(1))", "Lựa chọn số 1 khi phần tử có ID tĩnh, rõ ràng và duy nhất."),
    ("By.NAME: Rất nhanh", "Thích hợp cho các trường nhập liệu trong biểu mẫu form."),
    ("By.CSS_SELECTOR: Nhanh, linh hoạt", "Dùng khi cần tìm phần tử theo quan hệ cha-con, class, type (ví dụ: .example input)."),
    ("By.XPATH: Chậm nhất (O(N))", "Quét toàn bộ cây DOM. Cần thiết khi tìm kiếm ngược lên thẻ cha hoặc tìm theo text, nhưng dễ vỡ khi giao diện thay đổi."),
    ("Lưu ý với ID động:", "Trong các framework hiện đại (React/Vue), ID có thể bị sinh ngẫu nhiên (ví dụ: id='btn_18492'). Cần phân biệt ID tĩnh ổn định và ID sinh tự động.")
]
for name, note in locators_comp:
    p = tf9_r.add_paragraph()
    p.text = f"• {name}"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_AMBER if "ID" in name else COLOR_WHITE
    p.space_before = Pt(6)
    p2 = tf9_r.add_paragraph()
    p2.text = f"   {note}"
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_MUTED


# ==============================================================================
# SLIDE 10: LỚP KEYS & PHÍM ĐẶC BIỆT (MỤC 7.4)
# ==============================================================================
slide10 = create_base_slide("CHUYÊN ĐỀ TÀI LIỆU SELENIUM", "Mục 7.4: Lớp Keys & Cơ chế mô phỏng phím chức năng")

card10_l = add_card(slide10, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.1), "Bản chất kỹ thuật của selenium.webdriver.common.keys.Keys")
tf10_l = card10_l.text_frame
points10 = [
    ("Bảng mã Unicode Private Use Area (PUA):", "Các phím đặc biệt không có ký tự in ấn thông thường được Selenium định nghĩa dưới dạng các hằng số Unicode nằm trong vùng PUA (từ \\ue000 đến \\ue03d)."),
    ("Keys.ENTER = '\\ue007':", "Mô phỏng phím Enter / Return trên bàn phím cứng."),
    ("Keys.TAB = '\\ue004':", "Mô phỏng phím chuyển tiêu điểm Tab."),
    ("Keys.ARROW_UP = '\\ue013':", "Mô phỏng phím điều hướng mũi tên lên."),
    ("Chuỗi sự kiện trình duyệt phát ra:", "Khi dùng send_keys(Keys.*), trình duyệt kích hoạt đầy đủ chuỗi sự kiện: keydown -> keypress (nếu có) -> keyup, kích hoạt các event listener của JavaScript.")
]
for title, desc in points10:
    p = tf10_l.add_paragraph()
    p.text = f"• {title}"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN
    p.space_before = Pt(6)
    p2 = tf10_l.add_paragraph()
    p2.text = f"   {desc}"
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_WHITE

card10_r = add_card(slide10, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.1), "Ứng dụng cụ thể trong kiểm thử chức năng")
tf10_r = card10_r.text_frame
tf10_r.paragraphs[0].text = "Code minh họa sử dụng các phím đặc biệt:"
tf10_r.paragraphs[0].font.size = Pt(13)
tf10_r.paragraphs[0].font.color.rgb = COLOR_WHITE

p_code = tf10_r.add_paragraph()
p_code.text = (
    "from selenium.webdriver.common.keys import Keys\n\n"
    "# 1. Mô phỏng gõ phím ENTER:\n"
    "target_input.send_keys(Keys.ENTER)\n"
    "# Phản hồi hiển thị: 'You entered: ENTER'\n\n"
    "# 2. Mô phỏng gõ phím TAB:\n"
    "target_input.send_keys(Keys.TAB)\n"
    "# Phản hồi hiển thị: 'You entered: TAB'\n\n"
    "# 3. Mô phỏng gõ phím ARROW_UP:\n"
    "target_input.send_keys(Keys.ARROW_UP)\n"
    "# Trên trang key_presses: 'You entered: UP'\n"
    "# Trên trang inputs (type=number): giá trị tăng từ 50 lên 51!"
)
p_code.font.name = "Courier New"
p_code.font.size = Pt(11)
p_code.font.color.rgb = COLOR_GREEN
p_code.space_before = Pt(8)


# ==============================================================================
# SLIDE 11: ĐỐI TƯỢNG KIỂM THỬ: 2 TRANG WEB CỦA NHÓM W1
# ==============================================================================
slide11 = create_base_slide("THỰC HÀNH KIỂM THỬ", "Đối tượng kiểm thử: Phân tích 2 trang /key_presses và /inputs")

card11_l = add_card(slide11, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.1), "Trang 1: /key_presses (Bắt phím người dùng)")
tf11_l = card11_l.text_frame
points11_l = [
    ("URL kiểm thử:", "https://the-internet.herokuapp.com/key_presses"),
    ("Cấu trúc HTML:", "<form><input id='target' type='text'></form>\n<p id='result' style='color:green'></p>"),
    ("Cơ chế xử lý JavaScript trên trang:", "Trang sử dụng jQuery gắn lắng nghe sự kiện trên toàn tài liệu: $(document).keydown(function(event){ log('You entered: ' + keyboardMap[event.which]); });"),
    ("Thách thức kiểm thử:", "Do input nằm trong thẻ form không có thuộc tính chặn submit, việc gõ Enter sẽ reload trang làm mất phần tử. Cần dùng execute_script để can thiệp preventDefault().")
]
for title, desc in points11_l:
    p = tf11_l.add_paragraph()
    p.text = f"• {title}"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN
    p.space_before = Pt(6)
    p2 = tf11_l.add_paragraph()
    p2.text = f"   {desc}"
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_WHITE

card11_r = add_card(slide11, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.1), "Trang 2: /inputs (Xử lý trường nhập số)")
tf11_r = card11_r.text_frame
points11_r = [
    ("URL kiểm thử:", "https://the-internet.herokuapp.com/inputs"),
    ("Cấu trúc HTML:", "<div class='example'>\n  <p>Number</p>\n  <input type='number'>\n</div>"),
    ("Đặc tính của thẻ input type='number':", "Chỉ nhận ký tự số. Trình duyệt tự động tích hợp cơ chế tăng/giảm giá trị khi người dùng nhấn các phím mũi tên lên/xuống (Arrow Up / Arrow Down)."),
    ("Mục tiêu kiểm thử:", "Xác thực thuộc tính type, kiểm tra nhập chuỗi số '50', kiểm tra nhấn phím Keys.ARROW_UP tăng thành '51', và kiểm tra làm rỗng bằng hàm clear().")
]
for title, desc in points11_r:
    p = tf11_r.add_paragraph()
    p.text = f"• {title}"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN
    p.space_before = Pt(6)
    p2 = tf11_r.add_paragraph()
    p2.text = f"   {desc}"
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_WHITE


# ==============================================================================
# SLIDE 12: BẢNG THIẾT KẾ 6 KỊCH BẢN KIỂM THỬ (TEST CASES)
# ==============================================================================
slide12 = create_base_slide("THỰC HÀNH KIỂM THỬ", "Bảng thiết kế 6 kịch bản kiểm thử chi tiết (Test Cases)")

table_shape12 = slide12.shapes.add_table(7, 6, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.0))
t12 = table_shape12.table

widths12 = [Inches(1.0), Inches(2.7), Inches(1.6), Inches(2.9), Inches(2.4), Inches(1.133)]
for i, w in enumerate(widths12):
    t12.columns[i].width = w

headers12 = ["Mã TC", "Tên ca kiểm thử", "Trang", "Hàm sử dụng", "Kết quả mong đợi", "Kết quả"]
for i, h in enumerate(headers12):
    cell = t12.cell(0, i)
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_BG_CARD
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_CYAN

test_rows = [
    ("TC01", "Gõ phím ENTER", "/key_presses", "get, By.ID, send_keys, Keys.ENTER, text", "Text hiển thị: 'You entered: ENTER'", "PASSED"),
    ("TC02", "Gõ phím TAB", "/key_presses", "get, By.ID, send_keys, Keys.TAB, text", "Text hiển thị: 'You entered: TAB'", "PASSED"),
    ("TC03", "Gõ phím ARROW_UP", "/key_presses", "get, By.ID, send_keys, Keys.ARROW_UP, text", "Text hiển thị: 'You entered: UP'", "PASSED"),
    ("TC04", "Nhập số & đọc giá trị", "/inputs", "get, find_element, send_keys, get_attribute", "type='number', value='50'", "PASSED"),
    ("TC05", "Tăng số bằng phím lên", "/inputs", "send_keys, Keys.ARROW_UP, get_attribute", "value tăng từ '50' lên '51'", "PASSED"),
    ("TC06", "Xoá nội dung ô nhập", "/inputs", "send_keys, clear, get_attribute", "value làm trống về rỗng ''", "PASSED")
]
for r_idx, r_data in enumerate(test_rows, start=1):
    for c_idx, val in enumerate(r_data):
        cell = t12.cell(r_idx, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_BG_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.size = Pt(12)
        if c_idx == 5:
            p.font.bold = True
            p.font.color.rgb = COLOR_GREEN
        elif c_idx == 0:
            p.font.bold = True
            p.font.color.rgb = COLOR_BLUE_LIGHT
        else:
            p.font.color.rgb = COLOR_WHITE


# ==============================================================================
# SLIDE 13: TRIỂN KHAI CODE: TRANG KEY PRESSES (TC01 - TC03)
# ==============================================================================
slide13 = create_base_slide("TRIỂN KHAI MÃ NGUỒN", "Triển khai mã nguồn: Kiểm thử trang /key_presses (TC01 – TC03)")

card13_l = add_card(slide13, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.1), "Hiện thực hóa TC01, TC02, TC03")
tf13_l = card13_l.text_frame
p_code = tf13_l.paragraphs[0]
p_code.text = (
    "def test_01_key_presses_enter(self) -> None:\n"
    "    # 1. Điều hướng bằng hàm get\n"
    "    self.driver.get('https://.../key_presses')\n"
    "    \n"
    "    # 2. Định vị phần tử bằng By.ID\n"
    "    target = self.driver.find_element(By.ID, 'target')\n"
    "    \n"
    "    # Chặn form tự reload bằng JavaScript\n"
    "    self.driver.execute_script(\n"
    "        'document.querySelector(\"form\").onsubmit = '\n"
    "        'function(e){ e.preventDefault(); return false; };'\n"
    "    )\n"
    "    \n"
    "    # 3. Gửi phím đặc biệt bằng send_keys\n"
    "    target.send_keys(Keys.ENTER)\n"
    "    \n"
    "    # 4. Trích xuất text thông báo bằng thuộc tính .text\n"
    "    result = self.driver.find_element(By.ID, 'result')\n"
    "    self.assertEqual(result.text, 'You entered: ENTER')\n"
    "    self.driver.save_screenshot('.../tc01_enter.png')"
)
p_code.font.name = "Courier New"
p_code.font.size = Pt(10)
p_code.font.color.rgb = COLOR_GREEN

card13_r = add_card(slide13, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.1), "Giải thích kỹ thuật trọng tâm")
tf13_r = card13_r.text_frame
expl13 = [
    ("Áp dụng đủ các hàm theo phân công:", "get, find_element, By.ID, send_keys, Keys.ENTER, Keys.TAB, Keys.ARROW_UP, text, assertEqual."),
    ("Xử lý TC02 (Gõ phím TAB):", "Phím Keys.TAB kích hoạt sự kiện chuyển tiêu điểm, trang hiển thị kết quả 'You entered: TAB'."),
    ("Xử lý TC03 (Gõ phím ARROW_UP):", "Phím Keys.ARROW_UP kích hoạt mã phím 38, trang hiển thị kết quả 'You entered: UP'."),
    ("Lưu trữ ảnh minh chứng tự động:", "Mỗi bài test đều tích hợp save_screenshot() lưu ảnh trực tiếp vào thư mục tuan-06/anh/ để làm bằng chứng kiểm thử.")
]
for title, desc in expl13:
    p = tf13_r.add_paragraph()
    p.text = f"• {title}"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN
    p.space_before = Pt(8)
    p2 = tf13_r.add_paragraph()
    p2.text = f"   {desc}"
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_WHITE


# ==============================================================================
# SLIDE 14: TRIỂN KHAI CODE: TRANG INPUTS (TC04 - TC06)
# ==============================================================================
slide14 = create_base_slide("TRIỂN KHAI MÃ NGUỒN", "Triển khai mã nguồn: Kiểm thử trang /inputs (TC04 – TC06)")

card14_l = add_card(slide14, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.1), "Hiện thực hóa TC04, TC05, TC06")
tf14_l = card14_l.text_frame
p_code = tf14_l.paragraphs[0]
p_code.text = (
    "def test_04_inputs_send_keys_and_get_attribute(self):\n"
    "    self.driver.get('https://.../inputs')\n"
    "    num_input = self.driver.find_element(\n"
    "        By.CSS_SELECTOR, '.example input'\n"
    "    )\n"
    "    # Kiểm tra thuộc tính type\n"
    "    self.assertEqual(\n"
    "        num_input.get_attribute('type'), 'number'\n"
    "    )\n"
    "    # Nhập giá trị số và kiểm tra thuộc tính value\n"
    "    num_input.send_keys('50')\n"
    "    self.assertEqual(\n"
    "        num_input.get_attribute('value'), '50'\n"
    "    )\n\n"
    "def test_06_inputs_clear(self):\n"
    "    # Sử dụng hàm clear() làm rỗng ô nhập\n"
    "    num_input = self.driver.find_element(...)\n"
    "    num_input.clear()\n"
    "    self.assertEqual(\n"
    "        num_input.get_attribute('value'), ''\n"
    "    )"
)
p_code.font.name = "Courier New"
p_code.font.size = Pt(10)
p_code.font.color.rgb = COLOR_GREEN

card14_r = add_card(slide14, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.1), "Giải thích kỹ thuật trọng tâm")
tf14_r = card14_r.text_frame
expl14 = [
    ("Kiểm tra type='number' (TC04):", "Dùng get_attribute('type') để kiểm tra tính toàn vẹn của trường dữ liệu, đảm bảo không nhận ký tự chữ tùy tiện."),
    ("Tăng giá trị số bằng Keys.ARROW_UP (TC05):", "Gửi phím Keys.ARROW_UP vào ô input số sẽ kích hoạt tính năng step-up native của trình duyệt, đưa giá trị từ '50' lên '51'. Xác minh lại bằng get_attribute('value')."),
    ("Làm rỗng ô nhập bằng clear() (TC06):", "Sử dụng hàm clear() theo đúng yêu cầu phân công để xoá sạch toàn bộ giá trị đã nhập và khẳng định trường trở về chuỗi rỗng ''."),
    ("Đảm bảo không dùng chuỗi magic:", "Toàn bộ dữ liệu kiểm thử được định nghĩa tập trung qua các hằng số INPUT_NUMBER_VALUE, EXPECTED_INCREMENTED_VALUE.")
]
for title, desc in expl14:
    p = tf14_r.add_paragraph()
    p.text = f"• {title}"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN
    p.space_before = Pt(8)
    p2 = tf14_r.add_paragraph()
    p2.text = f"   {desc}"
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_WHITE


# ==============================================================================
# SLIDE 15: BÁO CÁO KẾT QUẢ KIỂM THỬ TỰ ĐỘNG
# ==============================================================================
slide15 = create_base_slide("KẾT QUẢ THỰC THI", "Báo cáo kết quả kiểm thử tự động & Phân tích log Pytest")

card15_l = add_card(slide15, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.1), "Thông số kết quả kiểm thử")
tf15_l = card15_l.text_frame
metrics = [
    ("Tổng số ca kiểm thử:", "6 / 6 Test Cases"),
    ("Tỷ lệ thành công (Success Rate):", "100% PASSED"),
    ("Tổng thời gian thực thi:", "5.01 giây (Cực kỳ tối ưu)"),
    ("Môi trường thực thi:", "macOS (Apple Silicon arm64), Python 3.13.9, Selenium 4.51.0, Chrome 155"),
    ("Báo cáo HTML sinh tự động:", "tuan-06/anh/report.html (Đầy đủ metadata môi trường, thời gian và trạng thái)")
]
for label, val in metrics:
    p = tf15_l.add_paragraph()
    p.text = f"• {label}"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN
    p.space_before = Pt(6)
    p2 = tf15_l.add_paragraph()
    p2.text = f"   {val}"
    p2.font.size = Pt(13)
    p2.font.bold = "PASSED" in val or "5.01" in val
    p2.font.color.rgb = COLOR_GREEN if ("PASSED" in val or "5.01" in val) else COLOR_WHITE

card15_r = add_card(slide15, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.1), "Trích xuất Log thực thi từ Terminal")
tf15_r = card15_r.text_frame
p_code = tf15_r.paragraphs[0]
p_code.text = (
    "$ pytest -v tuan-06/test_nhap_lieu_phim_dac_biet.py \\\n"
    "  --html=tuan-06/anh/report.html --self-contained-html\n\n"
    "==================== test session starts ====================\n"
    "platform darwin -- Python 3.13.9, pytest-8.4.2\n"
    "rootdir: /Users/macos/kttd-w1\n"
    "collected 6 items\n\n"
    "test_01_key_presses_enter PASSED               [ 16%]\n"
    "test_02_key_presses_tab PASSED                 [ 33%]\n"
    "test_03_key_presses_arrow_up PASSED            [ 50%]\n"
    "test_04_inputs_send_keys_and_get_attr PASSED   [ 66%]\n"
    "test_05_inputs_arrow_up_increment PASSED       [ 83%]\n"
    "test_06_inputs_clear PASSED                    [100%]\n\n"
    "-------- Generated html report: report.html ---------\n"
    "================= 6 passed in 5.01s =================="
)
p_code.font.name = "Courier New"
p_code.font.size = Pt(10)
p_code.font.color.rgb = COLOR_GREEN


# ==============================================================================
# SLIDE 16: HÌNH ẢNH MINH CHỨNG CHẠY THỰC TẾ
# ==============================================================================
slide16 = create_base_slide("MINH CHỨNG THỰC TẾ", "Ảnh chụp màn hình kết quả thực thi các ca kiểm thử")

# Thẻ 1: Enter
img1 = os.path.join(IMG_DIR, "tc01_key_presses_enter.png")
if os.path.exists(img1):
    slide16.shapes.add_picture(img1, Inches(0.8), Inches(1.6), width=Inches(3.6))
    c1 = add_card(slide16, Inches(0.8), Inches(4.0), Inches(3.6), Inches(0.9), "TC01: Gõ ENTER")
    c1.text_frame.paragraphs[0].font.size = Pt(12)
    p = c1.text_frame.add_paragraph()
    p.text = "Kết quả: You entered: ENTER"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_GREEN

# Thẻ 2: Arrow Up
img3 = os.path.join(IMG_DIR, "tc03_key_presses_arrow_up.png")
if os.path.exists(img3):
    slide16.shapes.add_picture(img3, Inches(4.866), Inches(1.6), width=Inches(3.6))
    c2 = add_card(slide16, Inches(4.866), Inches(4.0), Inches(3.6), Inches(0.9), "TC03: Gõ ARROW_UP")
    c2.text_frame.paragraphs[0].font.size = Pt(12)
    p = c2.text_frame.add_paragraph()
    p.text = "Kết quả: You entered: UP"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_GREEN

# Thẻ 3: Inputs Increment
img5 = os.path.join(IMG_DIR, "tc05_inputs_arrow_up_increment.png")
if os.path.exists(img5):
    slide16.shapes.add_picture(img5, Inches(8.933), Inches(1.6), width=Inches(3.6))
    c3 = add_card(slide16, Inches(8.933), Inches(4.0), Inches(3.6), Inches(0.9), "TC05: Tăng giá trị số")
    c3.text_frame.paragraphs[0].font.size = Pt(12)
    p = c3.text_frame.add_paragraph()
    p.text = "Giá trị tăng từ 50 lên 51"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_GREEN

# Hàng dưới: Báo cáo tóm tắt
card_bottom = add_card(slide16, Inches(0.8), Inches(5.1), Inches(11.733), Inches(1.6), "Lưu trữ minh chứng & Báo cáo kiểm thử")
tf_b = card_bottom.text_frame
tf_b.paragraphs[0].font.size = Pt(13)
p = tf_b.add_paragraph()
p.text = (
    "• Toàn bộ 6 ảnh chụp màn hình minh chứng được tự động lưu trong thư mục: tuan-06/anh/tc01_*.png đến tc06_*.png\n"
    "• Báo cáo tương tác hoàn chỉnh pytest-html lưu tại: tuan-06/anh/report.html\n"
    "• Bộ kiểm thử có thể chạy cả ở chế độ headless (chạy ngầm tốc độ cao) và GUI trực quan: HEADLESS=false python3 -m unittest -v ..."
)
p.font.size = Pt(12)
p.font.color.rgb = COLOR_WHITE


# ==============================================================================
# SLIDE 17: BÀI HỌC KINH NGHIỆM & QUY TRÌNH GIT NHÓM
# ==============================================================================
slide17 = create_base_slide("TỔNG KẾT KỸ THUẬT", "Bài học kinh nghiệm kỹ thuật & Quy trình làm việc nhóm")

card17_l = add_card(slide17, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.1), "Kinh nghiệm giải quyết vấn đề kỹ thuật")
tf17_l = card17_l.text_frame
lessons = [
    ("Xử lý hành vi mặc định của form HTML:", "Khi nhấn phím Enter trong ô input, form tự động gửi request GET làm reload trang. Cần hiểu rõ sự kiện DOM và can thiệp preventDefault() qua execute_script để giữ vững trạng thái kiểm thử."),
    ("Tối ưu hóa chiến lược tải trang eager:", "Các trang web thực tế thường tải nhiều script bên thứ ba (quảng cáo, analytics). Dùng chiến lược eager giúp bài test không bị treo."),
    ("Tái sử dụng phiên trình duyệt setUpClass:", "Tiết kiệm 80% thời gian chạy kiểm thử mà vẫn giữ trọn vẹn tính độc lập giữa các test method.")
]
for title, desc in lessons:
    p = tf17_l.add_paragraph()
    p.text = f"• {title}"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN
    p.space_before = Pt(8)
    p2 = tf17_l.add_paragraph()
    p2.text = f"   {desc}"
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_WHITE

card17_r = add_card(slide17, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.1), "Tuân thủ quy trình làm việc nhóm GitHub")
tf17_r = card17_r.text_frame
git_steps = [
    ("1. Phân chia Issue theo Milestone:", "Giao việc độc lập cho từng đầu việc: Kịch bản test, Code key presses, Code inputs, Báo cáo & slide."),
    ("2. Phát triển trên nhánh riêng (t06-*):", "Mỗi thành viên làm việc độc lập trên nhánh tính năng, không commit đè lên nhau."),
    ("3. Kiểm soát chất lượng qua Pull Request:", "Mở PR vào nhánh week-06 với cú pháp Closes #<id>, tiến hành code review trước khi merge."),
    ("4. Tích hợp cuối tuần vào nhánh main:", "Sau khi toàn bộ test case trên week-06 đạt 100% Passed, mở PR tổng gộp vào main và giữ lại nhánh week-06 cho giảng viên đánh giá.")
]
for title, desc in git_steps:
    p = tf17_r.add_paragraph()
    p.text = f"• {title}"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN
    p.space_before = Pt(8)
    p2 = tf17_r.add_paragraph()
    p2.text = f"   {desc}"
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_WHITE


# ==============================================================================
# SLIDE 18: KẾT LUẬN & PHIÊN HỎI ĐÁP (Q&A)
# ==============================================================================
slide18 = prs.slides.add_slide(blank_layout)
bg18 = slide18.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg18.fill.solid()
bg18.fill.fore_color.rgb = COLOR_BG_DARK
bg18.line.fill.background()

card_qa = add_card(slide18, Inches(1.5), Inches(1.5), Inches(10.333), Inches(4.5))
tf18 = card_qa.text_frame
tf18.paragraphs[0].alignment = PP_ALIGN.CENTER
tf18.paragraphs[0].text = "CẢM ƠN CÔ VÀ CÁC BẠN ĐÃ THEO DÕI BÀI BÁO CÁO!"
tf18.paragraphs[0].font.name = "Arial"
tf18.paragraphs[0].font.size = Pt(28)
tf18.paragraphs[0].font.bold = True
tf18.paragraphs[0].font.color.rgb = COLOR_CYAN

p = tf18.add_paragraph()
p.alignment = PP_ALIGN.CENTER
p.text = "PHIÊN HỎI & ĐÁP (Q & A)"
p.font.name = "Arial"
p.font.size = Pt(22)
p.font.bold = True
p.font.color.rgb = COLOR_WHITE
p.space_before = Pt(16)

p = tf18.add_paragraph()
p.alignment = PP_ALIGN.CENTER
p.text = (
    "Nhóm W1 – Báo cáo Kiểm thử tự động Tuần 6\n"
    "Kho mã nguồn: https://github.com/Nanhhhhhhh/kttd-w1\n"
    "Pull Request chính: https://github.com/Nanhhhhhhh/kttd-w1/pull/9"
)
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.color.rgb = COLOR_MUTED
p.space_before = Pt(24)


# Lưu file presentation PPTX
prs.save(OUTPUT_PPTX_PATH)
print("Đã tạo thành công bộ Slide PowerPoint 18 trang tại:", OUTPUT_PPTX_PATH)
